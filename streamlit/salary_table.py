"""
salary_table.py

Reads data/salary_table.xlsx -- the salary spreadsheet filled in by hand: one row per player, one column per season
(1996-97 ... 2033-34), and a "<season> status" column per season on the same row (Player option, Team option, UFA...).

read_table()        -> every filled-in cell as a long table: ROW, PLAYER_ID, PLAYER, TEAM, SEASON, SALARY, STATUS
                       (the Salaries sheet, plus the Dead money sheet's rows with STATUS "Dead money")
to_salary_data(t)   -> what the rest of the app reads: PLAYER_NAME, SEASON, SALARY, STATUS, TEAM, PLAYER_ID,
                       YEARS_REMAINING -- one row per player per season, repeated under each spelling of his name
                       (the sheet's, the NBA's official accented one, and an accent-free one) so every page's exact-name
                       lookups find him. Dead money is left out (it isn't the player's salary). Anything the sheet
                       doesn't have yet falls back to the old data/salaries.csv.
team_contracts(...) -> one team's rows for Front Office > Salary Cap (players found through the team's NBA roster).

The sheet is read by its headers, not fixed column letters, so adding, moving or re-sorting columns and rows is fine.
"""

import os
import re
import unicodedata

import pandas as pd

_HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.normpath(os.path.join(_HERE, "..", "data", "salary_table.xlsx"))
CSV = os.path.normpath(os.path.join(_HERE, "..", "data", "salaries.csv"))

STATUSES = ["Guaranteed", "Player option", "Team option", "Mutual option", "Non-guaranteed", "Partially guaranteed",
            "Two-way", "Estimate", "Dead money", "UFA", "RFA"]
# (background, text) -- the same colours the spreadsheet uses
STATUS_COLORS = {
    "Player option": ("#8CC152", "#000000"), "Team option": ("#F15A5A", "#FFFFFF"),
    "Mutual option": ("#966800", "#FFFFFF"),
    "Non-guaranteed": ("#D9D9D9", "#595959"), "Partially guaranteed": ("#FFE699", "#7F6000"),
    "Two-way": ("#BDD7EE", "#000000"), "Estimate": ("#B800FF", "#FFFFFF"), "Dead money": ("#B4A7D6", "#000000"),
    "UFA": ("#0B84D8", "#FFFFFF"), "RFA": ("#8B0000", "#FFFFFF"),
}
_STATUS_ALIASES = {
    "guaranteed": "Guaranteed", "g": "Guaranteed",
    "player option": "Player option", "po": "Player option", "player": "Player option",
    "team option": "Team option", "to": "Team option", "club option": "Team option", "team": "Team option",
    "mutual option": "Mutual option", "mo": "Mutual option", "mutual": "Mutual option",
    "estimate": "Estimate", "est": "Estimate", "estimated": "Estimate",
    "non guaranteed": "Non-guaranteed", "nonguaranteed": "Non-guaranteed", "ng": "Non-guaranteed",
    "partially guaranteed": "Partially guaranteed", "partial": "Partially guaranteed", "pg": "Partially guaranteed",
    "two way": "Two-way", "twoway": "Two-way", "2 way": "Two-way", "2w": "Two-way",
    "dead money": "Dead money", "dead": "Dead money", "dead cap": "Dead money", "waived": "Dead money",
    "ufa": "UFA", "rfa": "RFA",
}
_SEASON_RE = re.compile(r"^\s*(\d{4})\s*[-/]\s*(\d{2}|\d{4})\s*$")
_STATUS_RE = re.compile(r"^\s*(\d{4})\s*[-/]\s*(\d{2}|\d{4})\s+status\s*$", re.I)
_SUFFIX = {"jr", "sr", "ii", "iii", "iv", "v"}
TABLE_COLUMNS = ["ROW", "PLAYER_ID", "PLAYER", "TEAM", "SEASON", "SALARY", "STATUS"]
DATA_COLUMNS = ["PLAYER_NAME", "SEASON", "SALARY", "STATUS", "TEAM", "PLAYER_ID", "YEARS_REMAINING"]


def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")


def name_key(s):
    """Loose name match: no accents, case, dots/apostrophes, hyphens or Jr./III suffixes."""
    s = strip_accents(s).lower()
    s = re.sub(r"[.'`’]", "", s).replace("-", " ")
    return " ".join(t for t in s.split() if t not in _SUFFIX)


def _season(y1, y2):
    y1 = int(y1)
    return f"{y1}-{str(y1 + 1)[-2:]}"


def _money(v):
    """A typed salary -> whole dollars, or None. Takes numbers, '$58,456,566', '58.4M', '850K'. 0 stays 0 (a $0 cap
    hit, e.g. a training-camp deal); blanks and negatives are None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(round(v)) if v == v and v >= 0 else None
    s = str(v).strip().replace(",", "").replace("$", "").replace(" ", "")
    if not s or s in {"-", "—"}:
        return None
    mult = 1
    if s[-1:].lower() == "m":
        mult, s = 1_000_000, s[:-1]
    elif s[-1:].lower() == "k":
        mult, s = 1_000, s[:-1]
    if s.startswith("(") and s.endswith(")"):
        return None
    try:
        x = float(s) * mult
    except ValueError:
        return None
    return int(round(x)) if x >= 0 else None


def _status(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    k = re.sub(r"[-_]", " ", s.lower())
    k = " ".join(k.split())
    return _STATUS_ALIASES.get(k, s)


def _player_id(v):
    try:
        i = int(float(v))
        return i if i > 0 else None
    except (TypeError, ValueError):
        return None


def _clean(v):
    """None for blanks and pandas' NaN (pandas 3 stores a missing string as NaN), else the value."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return v


def _pick_sheet(wb):
    """The 'Salaries' sheet, or failing that the first sheet with a Player column."""
    if "Salaries" in wb.sheetnames:
        return wb["Salaries"]
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if any(str(h or "").strip().lower() in {"player", "player name", "name"} for h in first):
            return ws
    return None


def _read_sheet(ws, row_offset=0, dead_money=False):
    """One sheet's filled-in cells. dead_money: every amount on it is money a team (its Team column) still owes."""
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    col_player = col_team = col_id = None
    sal_cols, st_cols = {}, {}
    for i, h in enumerate(header):
        h = "" if h is None else str(h).strip()
        low = h.lower()
        m_sal, m_st = _SEASON_RE.match(h), _STATUS_RE.match(h)
        if m_st:
            st_cols[_season(*m_st.groups())] = i
        elif m_sal:
            sal_cols[_season(*m_sal.groups())] = i
        elif low in {"player", "player name", "name"} and col_player is None:
            col_player = i
        elif low in {"team", "tm"} and col_team is None:
            col_team = i
        elif "id" in low.split() or low in {"nba player id", "player id", "nba id"}:
            col_id = i if col_id is None else col_id
    if col_player is None or not sal_cols or (dead_money and col_team is None):
        return []
    out = []
    for r_i, row in enumerate(rows, start=2):
        if col_player >= len(row):
            continue
        name = row[col_player]
        name = " ".join(str(name).split()) if name is not None else ""
        if not name:
            continue
        team = row[col_team] if col_team is not None and col_team < len(row) else None
        team = str(team).strip().upper() if team not in (None, "") else None
        if dead_money and not team:
            continue
        pid = _player_id(row[col_id]) if col_id is not None and col_id < len(row) else None
        for season in sorted(set(sal_cols) | set(st_cols)):
            ci, si = sal_cols.get(season), st_cols.get(season)
            salary = _money(row[ci]) if ci is not None and ci < len(row) else None
            status = _status(row[si]) if si is not None and si < len(row) else None
            if dead_money:
                if salary is None:
                    continue
                status = "Dead money"
            if salary is None and status is None:
                continue
            out.append((row_offset + r_i, pid, name, team, season, salary, status))
    return out


def read_table(path=XLSX):
    """
    Every filled-in salary or status cell in the workbook, one row per player per season: the Salaries sheet, plus the
    Dead money sheet (STATUS "Dead money", TEAM = the team that still owes it; ROW numbers from 100000 up).
    """
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return pd.DataFrame(columns=TABLE_COLUMNS)
    try:
        ws = _pick_sheet(wb)
        out = _read_sheet(ws) if ws is not None else []
        if "Dead money" in wb.sheetnames:
            out += _read_sheet(wb["Dead money"], row_offset=100000, dead_money=True)
        return pd.DataFrame(out, columns=TABLE_COLUMNS)
    except Exception:
        return pd.DataFrame(columns=TABLE_COLUMNS)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _nba_players():
    try:
        from nba_api.stats.static import players
        return players.get_players()
    except Exception:
        return []


def to_salary_data(table, csv_path=CSV, nba_players=None):
    """The app-wide salary table (see module docstring). Never raises; empty (with the right columns) if nothing."""
    try:
        return _to_salary_data(table, csv_path, nba_players)
    except Exception:
        try:
            return _csv_only(csv_path, nba_players)
        except Exception:
            return pd.DataFrame(columns=DATA_COLUMNS)


def _index_players(nba_players):
    by_id, by_key = {}, {}
    for p in (nba_players if nba_players is not None else _nba_players()):
        by_id[p["id"]] = p
        by_key.setdefault(name_key(p["full_name"]), []).append(p)
    return by_id, by_key


def _official(pid, name, by_id, by_key):
    if pid and pid in by_id:
        return by_id[pid]
    pool = by_key.get(name_key(name), [])
    active = [p for p in pool if p.get("is_active")]
    pick = active or pool
    return pick[0] if len(pick) == 1 else None


def _variants(names):
    seen = []
    for n in names:
        for v in (n, strip_accents(n)):
            if v and v not in seen:
                seen.append(v)
    return seen


def _csv_frame(csv_path):
    try:
        df = pd.read_csv(csv_path)
    except Exception:
        return pd.DataFrame(columns=["PLAYER_NAME", "SEASON", "SALARY"])
    if not {"PLAYER_NAME", "SEASON", "SALARY"} <= set(df.columns):
        return pd.DataFrame(columns=["PLAYER_NAME", "SEASON", "SALARY"])
    df = df[["PLAYER_NAME", "SEASON", "SALARY"]].copy()
    df["SALARY"] = df["SALARY"].map(_money)
    df = df.dropna(subset=["PLAYER_NAME", "SEASON", "SALARY"])
    return df[df["SALARY"] > 0]


def _csv_only(csv_path, nba_players):
    return _to_salary_data(pd.DataFrame(columns=TABLE_COLUMNS), csv_path, nba_players)


def _to_salary_data(table, csv_path, nba_players):
    by_id, by_key = _index_players(nba_players)
    t = table.copy() if table is not None else pd.DataFrame(columns=TABLE_COLUMNS)
    t = t[t["SALARY"].notna() & (t["STATUS"] != "Dead money")]
    t = t[t["SALARY"] > 0]                  # a $0 cap hit isn't a salary for the Trade Machine / Salary filter

    # one identity per player: his NBA ID when known (typed, or matched by name), else his loose name
    records = {}          # ident -> {"names": [...], "pid": id, "seasons": {season: [salary, status, team]}}
    for row in t.itertuples(index=False):
        typed_id, status, team = _player_id(row.PLAYER_ID), _clean(row.STATUS), _clean(row.TEAM)
        off = _official(typed_id, row.PLAYER, by_id, by_key)
        pid = typed_id or (off["id"] if off else None)
        ident = ("id", pid) if pid else ("name", name_key(row.PLAYER))
        rec = records.setdefault(ident, {"names": [], "pid": pid, "seasons": {}})
        for n in (row.PLAYER, off["full_name"] if off else None):
            if n and n not in rec["names"]:
                rec["names"].append(n)
        cur = rec["seasons"].get(row.SEASON)
        if cur is None:
            rec["seasons"][row.SEASON] = [int(row.SALARY), status, team]
        else:                                   # two rows for the same season (e.g. signed twice): add them up
            cur[0] += int(row.SALARY)
            cur[1] = cur[1] or status
            cur[2] = cur[2] or team

    # the old CSV fills in whatever the sheet doesn't have
    have = {(k, s) for rec in records.values() for n in rec["names"] for k in [name_key(n)] for s in rec["seasons"]}
    for row in _csv_frame(csv_path).itertuples(index=False):
        k = name_key(row.PLAYER_NAME)
        if (k, row.SEASON) in have:
            continue
        off = _official(None, row.PLAYER_NAME, by_id, by_key)
        pid = off["id"] if off else None
        ident = ("id", pid) if pid else ("name", k)
        rec = records.setdefault(ident, {"names": [], "pid": pid, "seasons": {}})
        for n in (row.PLAYER_NAME, off["full_name"] if off else None):
            if n and n not in rec["names"]:
                rec["names"].append(n)
        rec["seasons"].setdefault(row.SEASON, [int(row.SALARY), None, None])
        have.add((k, row.SEASON))

    out = []
    for rec in records.values():
        seasons = sorted(rec["seasons"])
        for s in seasons:
            salary, status, team = rec["seasons"][s]
            left = sum(1 for x in seasons if x >= s)
            for n in _variants(rec["names"]):
                out.append((n, s, salary, status, team, rec["pid"], left))
    df = pd.DataFrame(out, columns=DATA_COLUMNS)
    if df.empty:
        return df
    df = df.drop_duplicates(subset=["PLAYER_NAME", "SEASON"], keep="first").reset_index(drop=True)
    df["SALARY"] = df["SALARY"].astype("int64")
    df["PLAYER_ID"] = df["PLAYER_ID"].astype("Int64")
    df["YEARS_REMAINING"] = df["YEARS_REMAINING"].astype("int64")
    return df


def team_contracts(table, team_abbr, seasons, roster_ids=None, roster_names=None):
    """
    One team's contracts for the Salary Cap page: a list of rows, each {"player", "player_id", "dead": bool, "cells": {season:
    (salary or None, status or None)}}, biggest first-season salary first and dead money last, plus per-season totals
    (dead money in, two-way contracts out -- they don't count against the cap). Only rows with something filled in
    for these seasons.

    Who's on the team comes from its NBA roster: roster_ids (NBA player IDs) match the sheet's NBA Player ID column;
    players typed in without an ID match roster_names by name. Without a roster, a Team column in the sheet (if there
    is one) is used instead. Dead money rows come from the Dead money sheet's Team column.
    """
    if table is None or table.empty:
        return [], {s: 0 for s in seasons}
    abbr = str(team_abbr).upper()
    ids = {i for i in (_player_id(x) for x in (roster_ids or [])) if i}
    keys = {name_key(n) for n in (roster_names or []) if n}

    def on_team(x):
        team = _clean(x.TEAM)
        if _clean(x.STATUS) == "Dead money":
            return team == abbr
        if not ids and not keys:
            return team == abbr
        pid = _player_id(x.PLAYER_ID)
        return (pid in ids) if pid else (name_key(x.PLAYER) in keys)

    t = table[table["SEASON"].isin(seasons)]
    t = t[[on_team(x) for x in t.itertuples(index=False)]] if not t.empty else t
    rows = []
    for (_r, player), g in t.groupby(["ROW", "PLAYER"], sort=False):
        cells = {s: (None, None) for s in seasons}
        for x in g.itertuples(index=False):
            sal = _clean(x.SALARY)
            cells[x.SEASON] = (None if sal is None else int(sal), _clean(x.STATUS))
        dead = any(st == "Dead money" for _v, st in cells.values())
        pid = next((_player_id(x.PLAYER_ID) for x in g.itertuples(index=False) if _player_id(x.PLAYER_ID)), None)
        rows.append({"player": player, "player_id": pid, "dead": dead, "cells": cells})
    first = seasons[0] if seasons else None
    rows.sort(key=lambda r: (r["dead"], -(r["cells"].get(first, (None, None))[0] or 0),
                             r["cells"].get(first, (None, None))[0] == 0,      # two-way deals before $0 cap hits
                             -sum(v or 0 for v, _s in r["cells"].values()), r["player"]))
    totals = {s: sum((r["cells"][s][0] or 0) for r in rows if r["cells"][s][1] != "Two-way") for s in seasons}
    return rows, totals


# The league's official cap figures by season (NBA announcements: 2026-27 on June 30, 2026 -- cap $164.961M, tax
# $200.428M, first apron $209.015M, second apron $221.686M). A season missing here shows its totals only.
CAP_FIGURES = {
    "2025-26": {"cap": 154_647_000, "tax": 187_895_000, "apron1": 195_945_000, "apron2": 207_824_000},
    "2026-27": {"cap": 164_961_000, "tax": 200_428_000, "apron1": 209_015_000, "apron2": 221_686_000},
}

_OPTION_CODES = {"Player option": "PO", "Team option": "TO", "Mutual option": "MO"}


def contract_index(table):
    """{("id", player_id) or ("name", loose name): {season: (salary or None, status or None)}} from read_table(),
    leaving out dead money -- built once, then looked up per player by contract_summary()."""
    idx = {}
    if table is None or table.empty:
        return idx
    for x in table.itertuples(index=False):
        status = _clean(x.STATUS)
        if status == "Dead money":
            continue
        sal = _clean(x.SALARY)
        cell = (None if sal is None else int(sal), status)
        pid = _player_id(x.PLAYER_ID)
        for k in ((("id", pid),) if pid else ()) + (("name", name_key(x.PLAYER)),):
            idx.setdefault(k, {}).setdefault(x.SEASON, cell)
    return idx


def contract_summary(index, season, player_id=None, name=None):
    """
    (salary in `season`, years left counting `season`, "PO"/"TO"/"MO" if the last of those years is a player / team /
    mutual option, else "") -- or None when there's no salary for that season. Years run on from `season` for as long
    as there is a salary each season (a UFA/RFA year ends it).
    """
    seasons = index.get(("id", _player_id(player_id))) if player_id else None
    if seasons is None and name:
        seasons = index.get(("name", name_key(name)))
    if not seasons or season not in seasons or seasons[season][0] is None:
        return None
    y = int(season[:4])
    years, last_status = 0, None
    while True:
        s = f"{y}-{str(y + 1)[-2:]}"
        cell = seasons.get(s)
        if not cell or cell[0] is None or cell[1] in ("UFA", "RFA"):
            break
        years, last_status = years + 1, cell[1]
        y += 1
    return seasons[season][0], max(years, 1), _OPTION_CODES.get(last_status, "")


def is_paid(index, season, player_id=None, name=None):
    """True when the player makes more than $0 in `season` on a standard contract -- False for a two-way deal, a $0 cap
    hit, or no salary at all (the players listed last on the Front Office roster and in the Trade Machine)."""
    seasons = index.get(("id", _player_id(player_id))) if player_id else None
    if seasons is None and name:
        seasons = index.get(("name", name_key(name)))
    cell = (seasons or {}).get(season)
    return bool(cell and cell[0] and cell[0] > 0 and cell[1] != "Two-way")


def format_contract(summary):
    """(58456566, 4, "PO") -> "$58.4m, 4 yrs PO" (millions to one decimal, not rounded up)."""
    if not summary:
        return ""
    sal, yrs, code = summary
    m = int(sal // 100_000) / 10
    return f"${m:.1f}m, {yrs} yr{'s' if yrs != 1 else ''}" + (f" {code}" if code else "")
